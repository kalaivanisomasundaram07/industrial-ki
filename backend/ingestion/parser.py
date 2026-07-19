import os
import sqlite3
import uuid
from datetime import datetime
from dotenv import load_dotenv

# Load env variables from root directory
dotenv_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "../..", ".env"))
load_dotenv(dotenv_path)

SQLITE_PATH = os.getenv("SQLITE_PATH", "./industrial_ki.db")
DATA_DIR = os.getenv("DATA_DIR", "./data/raw")

def get_sqlite_conn():
    conn = sqlite3.connect(SQLITE_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_sqlite_conn()
    cursor = conn.cursor()
    
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS documents (
        id TEXT PRIMARY KEY,
        filename TEXT NOT NULL,
        doc_type TEXT NOT NULL,
        raw_text TEXT NOT NULL,
        reliability_weight REAL NOT NULL,
        ingested_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    )
    """)
    
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS review_queue (
        id TEXT PRIMARY KEY,
        triple_json TEXT NOT NULL,
        evidence_snippet TEXT NOT NULL,
        doc_id TEXT NOT NULL,
        score REAL NOT NULL,
        status TEXT DEFAULT 'pending',
        FOREIGN KEY(doc_id) REFERENCES documents(id)
    )
    """)
    
    conn.commit()
    conn.close()

def detect_doc_type_and_weight(filename: str):
    name_lower = filename.lower()
    if any(k in name_lower for k in ["sop", "regulation", "compliance"]):
        return "sop", 1.0
    elif "manual" in name_lower:
        return "manual", 0.9
    elif "log" in name_lower:
        return "log", 0.7
    elif any(k in name_lower for k in ["inspection", "report"]):
        return "inspection", 0.6
    else:
        return "unknown", 0.5

def parse_pdf(filepath: str) -> str:
    """Parse PDF and return plain text (for backward compat)."""
    text, _ = parse_pdf_with_pages(filepath)
    return text

def parse_pdf_with_pages(filepath: str):
    """
    Parse a PDF and return:
      (full_text: str, page_map: list of (page_number, page_text))
    page_map allows downstream chunking to record accurate page numbers.
    """
    full_text = ""
    page_map = []

    try:
        import pdfplumber
        with pdfplumber.open(filepath) as pdf:
            for page_num, page in enumerate(pdf.pages, start=1):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    full_text += page_text + "\n"
                    page_map.append((page_num, page_text.strip()))
    except Exception as e:
        print(f"pdfplumber failed on {filepath}, trying fallback pypdf. Error: {e}")
        try:
            import pypdf
            reader = pypdf.PdfReader(filepath)
            for page_num, page in enumerate(reader.pages, start=1):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    full_text += page_text + "\n"
                    page_map.append((page_num, page_text.strip()))
        except Exception as e2:
            print(f"pypdf fallback failed on {filepath}. Error: {e2}")

    return full_text.strip(), page_map

def parse_docx(filepath: str):
    """
    Parse DOCX and return (text, page_map).
    Paragraphs are grouped into estimated pages (10 paragraphs per page).
    """
    PARAS_PER_PAGE = 10
    try:
        import docx
        doc = docx.Document(filepath)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        full_text = "\n".join(paragraphs)

        # Build page_map by grouping paragraphs
        page_map = []
        for i in range(0, len(paragraphs), PARAS_PER_PAGE):
            page_num = (i // PARAS_PER_PAGE) + 1
            page_text = "\n".join(paragraphs[i:i + PARAS_PER_PAGE])
            if page_text.strip():
                page_map.append((page_num, page_text))

        return full_text.strip(), page_map
    except Exception as e:
        print(f"docx parsing failed on {filepath}. Error: {e}")
        return "", []

def parse_xlsx(filepath: str):
    """
    Parse XLSX and return (text, page_map).
    Rows are grouped into estimated pages (50 rows per page).
    """
    ROWS_PER_PAGE = 50
    try:
        import openpyxl
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        sheet = wb.worksheets[0]
        rows_text = []
        for row in sheet.iter_rows(values_only=True):
            row_str = " ".join([str(cell) for cell in row if cell is not None])
            if row_str.strip():
                rows_text.append(row_str)
        full_text = "\n".join(rows_text)

        # Build page_map by grouping rows
        page_map = []
        for i in range(0, len(rows_text), ROWS_PER_PAGE):
            page_num = (i // ROWS_PER_PAGE) + 1
            page_text = "\n".join(rows_text[i:i + ROWS_PER_PAGE])
            if page_text.strip():
                page_map.append((page_num, page_text))

        return full_text.strip(), page_map
    except Exception as e:
        print(f"xlsx parsing failed on {filepath}. Error: {e}")
        return "", []

def parse_image(filepath: str) -> str:
    text = ""
    try:
        from unstructured.partition.image import partition_image
        elements = partition_image(filename=filepath)
        text = "\n".join([el.text for el in elements if hasattr(el, "text")])
    except Exception as e:
        print(f"unstructured.partition.image failed on {filepath}, trying direct pytesseract. Error: {e}")
        try:
            import pytesseract
            from PIL import Image
            img = Image.open(filepath)
            text = pytesseract.image_to_string(img)
        except Exception as e2:
            print(f"Direct pytesseract failed on {filepath}. Error: {e2}")
    return text.strip()

def parse_txt(filepath: str) -> str:
    try:
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            return f.read().strip()
    except Exception as e:
        print(f"txt parsing failed on {filepath}. Error: {e}")
        return ""

def parse_file(filepath: str):
    """
    Parse any supported file and return (text, page_map).
    page_map is a list of (page_number, page_text) tuples.
    """
    _, ext = os.path.splitext(filepath)
    ext = ext.lower()

    if ext == ".pdf":
        return parse_pdf_with_pages(filepath)
    elif ext in [".docx", ".doc"]:
        return parse_docx(filepath)
    elif ext in [".xlsx", ".xls"]:
        return parse_xlsx(filepath)
    elif ext in [".jpg", ".jpeg", ".png", ".bmp"]:
        raw_text = parse_image(filepath)
        page_map = [(1, raw_text)] if raw_text else []
        return raw_text, page_map
    elif ext == ".txt":
        raw_text = parse_txt(filepath)
        page_map = [(1, raw_text)] if raw_text else []
        return raw_text, page_map
    else:
        print(f"Unsupported file type for {filepath}")
        return "", []

def ingest_document(filepath: str) -> dict:
    filename = os.path.basename(filepath)
    doc_type, weight = detect_doc_type_and_weight(filename)
    raw_text, page_map = parse_file(filepath)

    if not raw_text:
        print(f"Warning: Extracted text is empty for {filepath}")
        return None

    doc_id = str(uuid.uuid4())
    conn = get_sqlite_conn()
    cursor = conn.cursor()

    cursor.execute("SELECT id FROM documents WHERE filename = ?", (filename,))
    row = cursor.fetchone()
    if row:
        existing_id = row["id"]
        cursor.execute("""
            UPDATE documents 
            SET doc_type = ?, raw_text = ?, reliability_weight = ?, ingested_at = CURRENT_TIMESTAMP
            WHERE id = ?
        """, (doc_type, raw_text, weight, existing_id))
        doc_id = existing_id
    else:
        cursor.execute("""
            INSERT INTO documents (id, filename, doc_type, raw_text, reliability_weight)
            VALUES (?, ?, ?, ?, ?)
        """, (doc_id, filename, doc_type, raw_text, weight))

    conn.commit()
    conn.close()

    return {
        "id": doc_id,
        "filename": filename,
        "doc_type": doc_type,
        "raw_text": raw_text,
        "page_map": page_map,
        "reliability_weight": weight,
    }

def scan_and_ingest() -> list:
    init_db()
    ingested_docs = []

    if not os.path.exists(DATA_DIR):
        os.makedirs(DATA_DIR, exist_ok=True)
        return ingested_docs

    for filename in os.listdir(DATA_DIR):
        filepath = os.path.join(DATA_DIR, filename)
        if os.path.isfile(filepath):
            doc_data = ingest_document(filepath)
            if doc_data:
                ingested_docs.append(doc_data)

    return ingested_docs
